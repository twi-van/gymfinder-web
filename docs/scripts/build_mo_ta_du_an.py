#!/usr/bin/env python3
"""Tạo DOCX mô tả dự án GYMFINDER (dùng upload Google Docs)."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path(__file__).resolve().parents[1] / "GymFinder_MoTaDuAn.docx"


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=16 if level == 1 else 13 if level == 2 else 12, bold=True)
    return p


def add_para(doc, text, *, bold=False, size=11, space_after=6, align=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run)
        p.paragraph_format.space_after = Pt(2)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        set_run_font(run, bold=True, size=10)
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(val))
            set_run_font(run, size=10)
    doc.add_paragraph()


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2)

    # Cover
    add_para(doc, "GYMFINDER", bold=True, size=22, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    add_para(
        doc,
        "MÔ TẢ DỰ ÁN HỆ THỐNG",
        bold=True,
        size=16,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=8,
    )
    add_para(
        doc,
        "Website tìm kiếm, đánh giá phòng tập và huấn luyện viên cá nhân tại TP. Hồ Chí Minh",
        size=12,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=4,
    )
    add_para(
        doc,
        "Phiên bản 1.0 • Cập nhật theo đặc tả và mã nguồn hiện tại (2026)",
        size=11,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=18,
    )
    add_para(
        doc,
        "Tài liệu này trình bày tổng quan nghiệp vụ, phạm vi, vai trò, chức năng, kiến trúc kỹ thuật và quy tắc vận hành của hệ thống GYMFINDER, phục vụ giảng viên / người chấm hiểu chi tiết hệ thống (có thể tải lên Google Docs).",
        size=11,
        space_after=12,
    )

    add_heading(doc, "1. Thông tin chung", 1)
    add_table(
        doc,
        ["Mục", "Nội dung"],
        [
            ["Tên dự án", "GYMFINDER (GymFinder)"],
            ["Loại sản phẩm", "Website tra cứu và đánh giá (không phải nền tảng đặt lịch / thanh toán)"],
            ["Phạm vi địa lý", "Phòng tập và huấn luyện viên cá nhân (PT) tại TP. Hồ Chí Minh"],
            ["Đối tượng sử dụng", "Guest (khách), User (thành viên), Admin (quản trị viên)"],
            ["Mục đích tài liệu", "Mô tả hệ thống đủ chi tiết để hiểu nghiệp vụ, chức năng và cách vận hành"],
            ["Tài liệu liên quan", "Đặc tả yêu cầu (DacTa), Thiết kế CSDL, Hướng dẫn chạy, README"],
        ],
    )

    add_heading(doc, "2. Bối cảnh và mục tiêu", 1)
    add_heading(doc, "2.1. Bối cảnh", 2)
    add_para(
        doc,
        "Người tập thường phải tìm thông tin phòng tập từ nhiều nguồn rời rạc. Giá, địa chỉ, tiện ích, loại hình và chất lượng dịch vụ khó so sánh. Đồng thời, nhu cầu tra cứu huấn luyện viên cá nhân (Personal Trainer – PT) theo chuyên môn, kinh nghiệm và đánh giá ngày càng phổ biến.",
    )
    add_heading(doc, "2.2. Mục tiêu", 2)
    add_bullets(
        doc,
        [
            "Tập trung thông tin phòng tập và PT tại TP.HCM trong một website.",
            "Hỗ trợ tìm kiếm, lọc và sắp xếp theo tiêu chí thực tế: khu vực, giá, đánh giá, loại hình, tiện ích.",
            "Cho phép lưu yêu thích và gửi đánh giá có kiểm duyệt.",
            "Cung cấp khu vực quản trị để duy trì dữ liệu và kiểm soát nội dung.",
            "Ưu tiên trải nghiệm desktop, đồng thời hỗ trợ responsive cơ bản trên màn hình nhỏ.",
        ],
    )
    add_heading(doc, "2.3. Định hướng sản phẩm", 2)
    add_para(
        doc,
        "GYMFINDER là website tra cứu và đánh giá. Người dùng tự liên hệ phòng tập/PT qua thông tin liên hệ được công bố trên hệ thống. Hệ thống không thay thế việc đặt lịch, thanh toán hay tư vấn trực tiếp.",
    )

    add_heading(doc, "3. Phạm vi hệ thống", 1)
    add_heading(doc, "3.1. Trong phạm vi", 2)
    add_bullets(
        doc,
        [
            "Đăng ký, đăng nhập, đăng xuất, quản lý hồ sơ cá nhân và đổi mật khẩu.",
            "Duyệt, tìm kiếm, lọc và sắp xếp phòng tập; xem chi tiết (mô tả, địa chỉ văn bản, ảnh, giá, giờ hoạt động, tiện ích, đánh giá).",
            "Tìm kiếm và xem hồ sơ PT; hỗ trợ nhiều chuyên môn và liên kết nhiều phòng tập theo thiết kế (UI sẵn sàng nhận nhiều liên kết).",
            "Quản lý danh sách yêu thích phòng tập và PT.",
            "Tạo / sửa / xóa đánh giá của chính người dùng; đánh giá mới phải được Admin duyệt trước khi công khai.",
            "Admin quản lý phòng tập, PT, danh mục (loại hình, tiện ích, chuyên môn, khu vực), người dùng và đánh giá; có dashboard tổng quan.",
            "Trang Điều khoản dịch vụ và Chính sách bảo mật; quên mật khẩu bằng OTP qua email — không SMS (3 bước UI + bảng password_reset_otps; gửi SMTP khi đủ hạ tầng).",
        ],
    )
    add_heading(doc, "3.2. Ngoài phạm vi", 2)
    add_bullets(
        doc,
        [
            "Không bản đồ, nút “Xem đường đi”, điều hướng hoặc liên kết Google Maps.",
            "Không đặt lịch PT/phòng tập, thanh toán, ví điện tử hoặc quản lý đơn hàng.",
            "Không chat thời gian thực, nhắn tin trong ứng dụng hoặc gọi điện tích hợp.",
            "Không ứng dụng mobile native; chỉ web responsive cơ bản.",
            "Không chatbot / AI phức tạp trong phạm vi mặc định.",
        ],
    )

    add_heading(doc, "4. Vai trò và phân quyền", 1)
    add_table(
        doc,
        ["Vai trò", "Quyền chính", "Giới hạn"],
        [
            [
                "Guest",
                "Xem nội dung công khai; tìm kiếm/lọc phòng tập & PT; xem chi tiết và đánh giá đã duyệt",
                "Không hồ sơ, yêu thích hay gửi đánh giá",
            ],
            [
                "User",
                "Toàn quyền Guest; cập nhật hồ sơ; quản lý yêu thích; tạo/sửa/xóa đánh giá của mình",
                "Không vào trang quản trị; không tự duyệt đánh giá",
            ],
            [
                "Admin",
                "Quản lý dữ liệu Gym/PT/danh mục; khóa/mở khóa tài khoản; kiểm duyệt đánh giá; xem dashboard",
                "Chỉ tài khoản role Admin được thao tác quản trị (Backend kiểm tra quyền, không chỉ ẩn nút UI)",
            ],
        ],
    )

    add_heading(doc, "5. Kiến trúc và công nghệ triển khai", 1)
    add_para(
        doc,
        "Hệ thống triển khai theo mô hình PHP Server-Side Rendering (SSR): trình duyệt nhận HTML đã render từ PHP, kết hợp Bootstrap cho giao diện. Một số thao tác động (ví dụ yêu thích) dùng API REST JSON.",
    )
    add_table(
        doc,
        ["Tầng", "Công nghệ"],
        [
            ["Giao diện", "PHP SSR, HTML5, CSS3, Bootstrap 5.3, JavaScript, Font Awesome 6, Google Fonts (Inter)"],
            ["Backend", "PHP 8.0+ thuần (không bắt buộc Composer)"],
            ["CSDL", "SQLite (mặc định khi chạy local) hoặc MySQL"],
            ["API", "Gateway api/ (favorites, reviews, …)"],
            ["Bảo mật", "PDO prepared statements, CSRF token, password_hash, phân quyền role, kiểm tra upload ảnh"],
        ],
    )
    add_para(doc, "Sơ đồ logic:", bold=True)
    add_para(
        doc,
        "Guest / User / Admin  →  User Portal (pages/)  &  Admin Portal (admin/)  →  PHP SSR + API  →  SQLite / MySQL",
        size=10,
    )
    add_heading(doc, "5.1. Cấu trúc thư mục chính", 2)
    add_bullets(
        doc,
        [
            "index.php, router.php — trang chủ và router PHP built-in server",
            "pages/ — trang người dùng (auth, gyms, trainers, user, legal)",
            "admin/pages/ — trang quản trị",
            "backend/ — bootstrap, auth, repositories, resources (nghiệp vụ)",
            "database/ — schema.sql, seed.sql, ERD",
            "includes/ — header, footer, card, layout dùng chung",
            "api/ — REST API",
            "DOCS/ — đặc tả, thiết kế CSDL, hướng dẫn chạy",
        ],
    )

    add_heading(doc, "6. Chức năng phía người dùng", 1)
    add_heading(doc, "6.1. Trang chủ", 2)
    add_bullets(
        doc,
        [
            "Giới thiệu mục đích GYMFINDER và tìm kiếm nhanh (từ khóa, khu vực, khoảng giá).",
            "Hiển thị phòng tập / HLV nổi bật khi được đánh dấu.",
            "Header, footer, logo và màu chủ đạo thống nhất toàn site.",
        ],
    )
    add_heading(doc, "6.2. Tài khoản và xác thực", 2)
    add_table(
        doc,
        ["Mã", "Chức năng", "Mô tả"],
        [
            ["AUTH-01", "Đăng ký", "Họ tên, email, mật khẩu (bắt buộc); SĐT tuỳ chọn; bắt buộc đồng ý Điều khoản & Chính sách"],
            ["AUTH-02", "Đăng nhập", "Email hoặc SĐT; hỗ trợ Ghi nhớ đăng nhập (cookie phiên ~30 ngày); báo lỗi khi sai hoặc tài khoản bị khóa"],
            ["AUTH-03", "Đăng xuất", "Kết thúc phiên, hết hạn cookie, chuyển trang phù hợp"],
            ["AUTH-04", "Phân quyền", "Backend kiểm tra quyền cho mọi thao tác nhạy cảm"],
            ["AUTH-05", "Tài khoản khóa", "User bị khóa không thực hiện thao tác cần quyền User"],
            ["AUTH-06", "Quên mật khẩu OTP", "3 bước: email → OTP 6 số (10 phút, ≤5 lần sai) → mật khẩu mới; lưu bảng password_reset_otps; SMTP bổ sung sau"],
        ],
    )
    add_heading(doc, "6.3. Hồ sơ cá nhân", 2)
    add_para(doc, "User xem/cập nhật họ tên, số điện thoại, ảnh đại diện và đổi mật khẩu (xác nhận mật khẩu mới, tối thiểu 8 ký tự).")
    add_heading(doc, "6.4. Tìm kiếm và chi tiết phòng tập", 2)
    add_bullets(
        doc,
        [
            "Lọc theo từ khóa, quận/khu vực, khoảng giá (4 mức chuẩn + khoảng tùy chỉnh), đánh giá (sao), loại hình, tiện ích, có PT.",
            "Sắp xếp: mới nhất, đánh giá cao nhất, giá tăng/giảm.",
            "Thẻ kết quả: tên, ảnh, địa chỉ văn bản, mức giá, điểm đánh giá.",
            "Chi tiết: mô tả, liên hệ, giờ hoạt động, gallery, tiện ích, HLV liên kết, đánh giá đã duyệt.",
            "Không hiển thị bản đồ / tọa độ / chỉ đường. Thiếu dữ liệu thì hiển thị “Chưa cập nhật” / “Chưa có đánh giá”, không bịa số liệu.",
        ],
    )
    add_heading(doc, "6.5. Tìm kiếm và hồ sơ PT", 2)
    add_bullets(
        doc,
        [
            "Danh sách PT; lọc theo từ khóa, khu vực, chuyên môn, phòng tập liên kết.",
            "Một PT có thể có nhiều chuyên môn (thiết kế N–N qua trainer_specialties); một PT có thể liên kết nhiều phòng tập (thiết kế N–N qua gym_trainers).",
            "Kết quả tìm kiếm không nhân bản một PT chỉ vì PT gắn nhiều phòng tập.",
            "Hồ sơ công khai chỉ liệt kê phòng tập đang hoạt động; không đặt lịch / thanh toán / chat.",
        ],
    )
    add_heading(doc, "6.6. Yêu thích", 2)
    add_bullets(
        doc,
        [
            "User thêm/bỏ yêu thích phòng tập hoặc PT; xem danh sách và mở trang chi tiết.",
            "Không trùng cùng một đối tượng; Guest bấm yêu thích thì được yêu cầu đăng nhập.",
        ],
    )
    add_heading(doc, "6.7. Đánh giá", 2)
    add_bullets(
        doc,
        [
            "User gửi đánh giá 1–5 sao + nội dung cho một phòng tập hoặc một PT.",
            "Mỗi User tối đa một đánh giá hiện hành cho mỗi đối tượng; có thể sửa (sau sửa quay lại chờ duyệt) hoặc xóa.",
            "Trạng thái: chờ duyệt (pending) → đã duyệt (approved) hoặc từ chối (rejected).",
            "Chỉ đánh giá đã duyệt được công khai và tính vào điểm trung bình / số lượt.",
            "Khi Admin từ chối: bắt buộc nhập lý do (1–255 ký tự); User có thể xem lý do.",
        ],
    )

    add_heading(doc, "7. Chức năng quản trị", 1)
    add_heading(doc, "7.1. Dashboard", 2)
    add_para(doc, "Hiển thị các chỉ số thực tế: số phòng tập đang công khai, PT đang hoạt động, người dùng, đánh giá chờ duyệt. Biểu đồ (nếu có) dựa trên dữ liệu thật.")
    add_heading(doc, "7.2. Quản lý phòng tập", 2)
    add_bullets(
        doc,
        [
            "Thêm, xem, tìm kiếm, sửa thông tin (tên, mô tả, địa chỉ văn bản, khu vực, liên hệ, giá, giờ, danh mục, tiện ích, ảnh).",
            "Quản lý hiển thị bằng Tạm ẩn / Kích hoạt lại — không có nút Xóa cứng trên giao diện quản trị.",
            "Phòng tập tạm ẩn biến mất khỏi tìm kiếm công khai, trang chủ, yêu thích; URL chi tiết trả 404; dữ liệu được giữ nguyên để kích hoạt lại.",
            "Không trường tọa độ / bản đồ trên UI quản trị.",
        ],
    )
    add_heading(doc, "7.3. Quản lý PT (HLV)", 2)
    add_bullets(
        doc,
        [
            "Thêm, sửa, tạm ẩn / kích hoạt lại HLV (không xóa cứng trên UI).",
            "Gán một hoặc nhiều chuyên môn; liên kết một hoặc nhiều phòng tập trên cùng hồ sơ (không nhân bản bản ghi PT).",
            "Mỗi PT trong phạm vi hiện tại phải liên kết ít nhất một phòng tập.",
        ],
    )
    add_heading(doc, "7.4. Danh mục", 2)
    add_para(doc, "Trang danh mục hợp nhất theo tab: loại phòng tập, chuyên môn HLV, tiện ích, khu vực (quận).")
    add_heading(doc, "7.5. Quản lý người dùng", 2)
    add_para(
        doc,
        "Xem danh sách; Khóa / Mở khóa tài khoản. Không có nút Xóa tài khoản trên UI. Khóa không làm mất lịch sử đánh giá hay yêu thích.",
    )
    add_heading(doc, "7.6. Kiểm duyệt đánh giá", 2)
    add_bullets(
        doc,
        [
            "Xem danh sách chờ duyệt; duyệt hoặc từ chối.",
            "Từ chối bắt buộc nhập lý do qua modal; Hủy không đổi trạng thái.",
            "Sau duyệt/từ chối, điểm trung bình và số lượng công khai được cập nhật nhất quán.",
        ],
    )
    add_heading(doc, "7.7. Hồ sơ Admin", 2)
    add_para(doc, "Admin chỉ cập nhật hồ sơ của chính mình; không tự đổi vai trò / tự cấp quyền qua màn hình hồ sơ.")

    add_heading(doc, "8. Quy tắc nghiệp vụ cốt lõi", 1)
    add_table(
        doc,
        ["Mã", "Quy tắc"],
        [
            ["BR-01", "Guest xem công khai; thao tác cá nhân yêu cầu đăng nhập."],
            ["BR-02", "Chỉ Admin thao tác quản trị."],
            ["BR-03", "Gym/PT tạm ẩn không xuất hiện công khai; URL chi tiết → 404."],
            ["BR-04", "PT–chuyên môn và Gym–Trainer theo quan hệ nhiều–nhiều; không PT tự do; không nhân bản hồ sơ khi gán thêm phòng tập."],
            ["BR-05", "Đánh giá phải được duyệt mới công khai và tính điểm."],
            ["BR-06", "User chỉ quản lý dữ liệu của chính mình."],
            ["BR-07", "Địa chỉ là văn bản; không bản đồ/chỉ đường."],
            ["BR-08", "Không đặt lịch, thanh toán, chat."],
            ["BR-09", "Không tự tạo giá/điểm/liên hệ giả khi thiếu dữ liệu."],
            ["BR-10", "Từ chối đánh giá bắt buộc có lý do."],
            ["BR-11", "Gym/HLV chỉ dùng Tạm ẩn / Kích hoạt lại trên UI (không xóa cứng)."],
            ["BR-12", "Khóa user không mất lịch sử; UI chỉ cập nhật khi Backend xác nhận thành công."],
        ],
    )

    add_heading(doc, "9. Trạng thái dữ liệu", 1)
    add_table(
        doc,
        ["Đối tượng", "Trạng thái / quy tắc hiển thị"],
        [
            ["Phòng tập", "active | hidden — chỉ active xuất hiện công khai"],
            ["PT", "active | hidden — tương tự; tạm ẩn vẫn giữ đánh giá lịch sử"],
            ["Đánh giá", "pending | approved | rejected — chỉ approved công khai"],
            ["Tài khoản", "active | locked — locked không đăng nhập được nhưng giữ lịch sử"],
            ["Giá / điểm", "Thiếu dữ liệu → “Chưa cập nhật” / “Chưa có đánh giá”"],
        ],
    )

    add_heading(doc, "10. Cơ sở dữ liệu (tóm tắt)", 1)
    add_para(
        doc,
        "Thiết kế chi tiết nằm trong tài liệu Thiết kế CSDL. Runtime hiện có các bảng chính trong schema (MySQL/SQLite):",
    )
    add_bullets(
        doc,
        [
            "users — tài khoản (role user/admin, status active/locked)",
            "districts, categories, amenities, specialties — danh mục dùng lọc/gán",
            "gyms, gym_images, gym_categories, gym_amenities — phòng tập và quan hệ phụ",
            "trainers — hồ sơ HLV (liên kết gym/chuyên môn theo triển khai hiện tại; đặc tả hướng tới bảng trung gian gym_trainers, trainer_specialties)",
            "reviews — đánh giá gym hoặc PT, có trạng thái và lý do từ chối",
            "favorites — yêu thích theo user",
            "password_reset_otps — OTP quên mật khẩu (hash, hết hạn, attempts, verified_at, consumed_at)",
        ],
    )
    add_para(
        doc,
        "Ghi chú cho giảng viên: đặc tả nghiệp vụ mô tả quan hệ nhiều–nhiều Gym–Trainer và Trainer–Specialty; giao diện quản trị/public đã sẵn sàng nhận nhiều liên kết. Quên mật khẩu dùng OTP email lưu bảng password_reset_otps; gửi SMTP có thể bổ sung. Chi tiết xem Thiết kế CSDL và database/schema.sql.",
        size=10,
    )

    add_heading(doc, "11. Bảo mật và yêu cầu phi chức năng", 1)
    add_bullets(
        doc,
        [
            "Mật khẩu băm an toàn (password_hash), không lưu dạng rõ.",
            "Truy vấn dùng PDO prepared statements.",
            "Escape dữ liệu khi hiển thị HTML (chống XSS).",
            "CSRF token cho thao tác thay đổi dữ liệu.",
            "Backend kiểm tra quyền Admin/User; không tin role từ client.",
            "Upload ảnh: giới hạn loại/kích thước, kiểm tra file, không thực thi file upload.",
            "Không lộ chuỗi kết nối / stack trace nhạy cảm ra giao diện.",
            "UI nhất quán; có trạng thái tải, lỗi, thành công, danh sách rỗng, modal xác nhận.",
        ],
    )

    add_heading(doc, "12. Các trang chính của hệ thống", 1)
    add_table(
        doc,
        ["Nhóm", "Trang", "Đường dẫn (ví dụ localhost:8000)"],
        [
            ["Công khai", "Trang chủ", "/"],
            ["Công khai", "Tìm phòng tập", "/pages/gyms/search.php"],
            ["Công khai", "Chi tiết phòng tập", "/pages/gyms/detail.php?id=1"],
            ["Công khai", "Danh sách HLV", "/pages/trainers/index.php"],
            ["Công khai", "Chi tiết HLV", "/pages/trainers/detail.php?id=1"],
            ["Công khai", "Điều khoản / Chính sách", "/pages/legal/terms.php , privacy.php"],
            ["Auth", "Đăng nhập / Đăng ký / Quên MK (OTP 3 bước)", "/pages/auth/login|register|forgot_password|forgot_verify_otp|reset_password.php"],
            ["User", "Yêu thích / Hồ sơ", "/pages/user/favorites.php , profile.php"],
            ["Admin", "Dashboard, Gyms, Trainers, Users, Reviews, Danh mục, Profile", "/admin/pages/*.php"],
        ],
    )

    add_heading(doc, "13. Cách chạy nhanh và tài khoản demo", 1)
    add_para(doc, "Khởi chạy (khuyến nghị PHP built-in server):", bold=True)
    add_para(doc, "php -S localhost:8000 router.php", size=10)
    add_para(doc, "Mở http://localhost:8000 — lần đầu tự tạo SQLite và nạp dữ liệu mẫu.")
    add_para(doc, "Tài khoản demo:", bold=True)
    add_table(
        doc,
        ["Vai trò", "Email", "Mật khẩu"],
        [
            ["Admin", "u01@gymfinder.test", "password"],
            ["User", "u02@gymfinder.test", "password"],
            ["Guest", "(không cần đăng nhập)", "—"],
        ],
    )

    add_heading(doc, "14. Phân loại ưu tiên triển khai", 1)
    add_table(
        doc,
        ["Mức", "Phạm vi"],
        [
            [
                "P0 – Bắt buộc",
                "Xác thực/phân quyền; tìm kiếm & chi tiết Gym/PT; hồ sơ; yêu thích; đánh giá/kiểm duyệt; quản trị cốt lõi; không bản đồ",
            ],
            ["P1 – Quan trọng", "Bộ lọc nâng cao; dashboard/biểu đồ; nổi bật; phòng tập tương tự; responsive tốt; hiện lý do từ chối cho user"],
            [
                "P2 – Có thể bổ sung",
                "Gửi OTP qua SMTP thật; trang pháp lý; UX/báo cáo bổ sung nếu đủ thời gian",
            ],
        ],
    )
    add_para(doc, "Không tự mở rộng sang đặt lịch, thanh toán, chat hoặc AI nếu chưa có quyết định mới.")

    add_heading(doc, "15. Thuật ngữ", 1)
    add_table(
        doc,
        ["Thuật ngữ", "Giải thích"],
        [
            ["Gym", "Phòng tập thể hình"],
            ["PT / HLV", "Personal Trainer – huấn luyện viên cá nhân"],
            ["Guest", "Khách chưa đăng nhập"],
            ["User", "Người dùng đã đăng nhập"],
            ["Admin", "Quản trị viên"],
            ["Tạm ẩn (hidden)", "Ngừng hiển thị công khai nhưng giữ dữ liệu để kích hoạt lại"],
            ["Kiểm duyệt", "Admin duyệt hoặc từ chối đánh giá trước khi công khai"],
            ["SSR", "Server-Side Rendering – PHP render HTML phía máy chủ"],
        ],
    )

    add_heading(doc, "16. Kết luận", 1)
    add_para(
        doc,
        "GYMFINDER là hệ thống web hỗ trợ người dùng tại TP.HCM tìm kiếm, so sánh và đánh giá phòng tập cùng huấn luyện viên cá nhân, kèm khu vực quản trị để duy trì dữ liệu và kiểm soát nội dung. Sản phẩm tập trung vào tra cứu minh bạch, không mở rộng sang đặt lịch hay thanh toán trong phạm vi hiện tại.",
    )
    add_para(
        doc,
        "Để xem chi tiết kỹ thuật hơn, vui lòng đối chiếu thêm: GYMFINDER_dacta.docx (đặc tả đầy đủ), GYMFINDER_CSDL.docx (CSDL), GymFinder_HuongDanChay.md (chạy hệ thống), và mã nguồn trên GitHub.",
        size=10,
    )

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
